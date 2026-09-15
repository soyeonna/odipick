#!/usr/bin/env python3
"""검수 요청 엑셀(작성완료본)을 index.html 에 반영한다 — 바뀐 칸만.

  ~/.agent-reach-venv/bin/python scripts/apply_review_xlsx.py docs/odipick_검수요청_v2_작성완료.xlsx [--dry]

- 상태가 '수정 확정' 인 줄만.
- 지금 사이트 값과 다른 칸만 반영한다 (내가 미리 채워 보낸 칸은 그대로라서 건너뜀).
- 반영 칸: 확정명, 주소, 업종, 영업시간, 주차, 룸, 대표메뉴, 메뉴·가격, 오디픽 소개글, 추천 이유, 주의점.
- 근거 없는 값은 채우지 않는다. 출처는 evidence 에 남긴다.
"""
import json, re, sys, datetime
from openpyxl import load_workbook
sys.path.insert(0, 'scripts')
XLSX=sys.argv[1]; DRY='--dry' in sys.argv; TODAY=datetime.date.today().isoformat()

def parse_prices(txt):
    out=[]
    for part in re.split(r'\s*/\s*', str(txt or '')):
        m=re.search(r'(.+?)\s*([\d,]{4,})\s*원', part)
        if not m: continue
        name=m.group(1).strip(); p=int(m.group(2).replace(',',''))
        t='groupMenu' if re.search(r'(\d)\s*[-–~]\s*(\d)\s*인', name) else \
          'side' if re.search(r'미나리|반찬|사리|볶음밥|공기밥', name) else \
          'set' if re.search(r'코스|정식|한정식|세트', name) else 'main'
        row={'m':name,'p':p,'type':t}
        gm=re.search(r'(\d)\s*[-–~]\s*(\d)\s*인', name)
        if gm: row['minPeople']=int(gm.group(1)); row['maxPeople']=int(gm.group(2))
        out.append(row)
    return out
def yn(v):
    v=str(v or '').strip()
    if v in ('있음','가능','있어요','O'): return True
    if v in ('없음','불가','없어요','X'): return False
    return None

H=open('index.html',encoding='utf-8').read()
m=re.search(r'(<script id="places" type="application/json">)\s*(\[.*?\])\s*(</script>)',H,re.S)
P=json.loads(m.group(2)); byname={p['n']:p for p in P}
wb=load_workbook(XLSX,data_only=True)
done=[]; held=[]; stats={}
for name in wb.sheetnames:
    if name=='작성기준': continue
    ws=wb[name]; hdr=[h for h in next(ws.iter_rows(min_row=1,max_row=1,values_only=True)) if h]
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not r or not r[0]: continue
        row=dict(zip(hdr,r[:len(hdr)])); s=lambda k: str(row.get(k) or '').strip()
        if s('상태')!='수정 확정': continue
        p=byname.get(s('기존명'))
        if not p: held.append((s('기존명'),'목록에 없음')); continue
        ch=[]
        if s('확정명') and s('확정명')!=p['n']: p['n']=s('확정명'); ch.append('이름')
        if s('주소') and s('주소')!=str(p.get('addr') or p.get('area') or ''): p['addr']=s('주소'); ch.append('주소')
        if s('업종') and s('업종')!=str(p.get('cat') or ''): p['cat']=s('업종'); ch.append('업종')
        if s('영업시간') and s('영업시간')!=str(p.get('hours') or ''): p['hours']=s('영업시간'); p['hoursCheckedAt']=TODAY; ch.append('영업시간')
        f=p.setdefault('fac',{})
        for col,key in (('주차','parking'),('룸','room')):
            v=yn(row.get(col))
            if v is not None and f.get(key)!=v: f[key]=v; ch.append(col)
        if s('대표메뉴') and s('대표메뉴')!=str(p.get('sig') or ''): p['sig']=s('대표메뉴'); ch.append('대표메뉴')
        pr=parse_prices(row.get('메뉴·가격'))
        if pr and pr!=p.get('prices'): p['prices']=pr; ch.append(f'메뉴 {len(pr)}개')
        for col,key in (('오디픽 소개글','v'),('추천 이유','why'),('주의점','caution')):
            if s(col) and s(col)!=str(p.get(key) or ''): p[key]=s(col); ch.append(col)
        if not ch: continue
        p['ownerCheckedAt']=TODAY
        p.setdefault('evidence',[]).append({'source':'대전공주 검수 엑셀 v2','fields':ch,'verifiedAt':TODAY,
                                            'refs':[u for u in s('출처').split() if u.startswith('http')][:3]})
        done.append((name,p['n'],', '.join(ch)))
        for c in ch: stats[c.split(' ')[0]]=stats.get(c.split(' ')[0],0)+1
print(f"반영 {len(done)}곳 — 칸별: {stats}")
for sh,n,c in done[:8]: print(f"  {sh} | {n} | {c}")
if held: print("보류:",held)
if DRY: sys.exit(0)
H=H[:m.start(2)]+json.dumps(P,ensure_ascii=False)+H[m.end(2):]
open('index.html','w',encoding='utf-8').write(H); print("index.html 반영")
